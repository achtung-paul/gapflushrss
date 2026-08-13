# Builds the blank .xlsx a user downloads before importing their own
# tolerance table (see core/tolerance_csv_import.py for the format this
# has to match, and CLAUDE.md for why this deliberately contains no real
# numbers -- a real industrial tolerance standard's actual figures were
# shared with this project for context only and must never be
# redistributed to end users).
#
# Generated on request rather than a static file checked into the repo:
# there's no way to hand-author a real .xlsx binary without a library,
# and generating it from the same understanding as the parser keeps the
# two from silently drifting apart if the expected format ever changes.

from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment

_INSTRUCTIONS = (
    "Row 1: leave columns A and B alone, then add one column per "
    'dimension range, labelled like "1 to 3" or "> 3 to 6" -- add as '
    "many as your standard needs.\n\n"
    "Each row below: column A is a tolerance group name, column B is "
    'T, NT, or blank (blank means the one value applies to both -- use '
    "this for a class your standard doesn't split by tooling).\n\n"
    "Leave a cell blank where your standard has no value for that "
    "group/range -- don't enter 0.\n\n"
    "When done, save as CSV (not .xlsx) and upload that file."
)


def build_import_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tolerance Data"

    sheet.append(["Group", "Type", "1 to 3", "> 3 to 6", "> 6 to 10"])
    # Example rows illustrating the T/NT pairing and the blank-type
    # style -- deliberately no values in any of them.
    sheet.append(["TG1", "T"])
    sheet.append(["TG1", "NT"])
    sheet.append(["TG2", ""])

    sheet["A1"].comment = Comment(_INSTRUCTIONS, "StackFit Pro")

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# The second supported layout (see core/tolerance_csv_import.py's module
# docstring) -- each nominal size range spans several rows, one per
# named accuracy grade, instead of the first template's T/NT pair.
_INSTRUCTIONS_V2 = (
    "Row 1: leave columns A and B alone, then add one column per "
    'dimension range, labelled like "from 0.1 to 1" or "over 1 up to 3" '
    "-- add as many as your standard needs. The last column may be "
    'labelled "over N" with no upper bound (e.g. "over 10") for an '
    "open-ended top range.\n\n"
    "Each block of rows below one nominal size range: column A holds "
    "the size range only on the block's first row -- leave it blank on "
    "the rows under it (they carry down automatically once you save as "
    "CSV). Column B is the accuracy grade -- it doesn't have to be "
    "exactly fine/medium/course/very course, use whatever your standard "
    "calls its grades.\n\n"
    "Leave a cell blank where your standard has no value for that "
    "range/grade -- don't enter 0.\n\n"
    "When done, save as CSV (not .xlsx) and upload that file."
)

_GRADE_EXAMPLES = ("fine", "medium", "course", "very course")


def build_import_template_v2() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tolerance Data"

    sheet.append(
        [
            "Nominal size range",
            "Accuracy grade",
            "from 0.1 to 1",
            "over 1 up to 3",
            "over 3 up to 6",
            "over 6 up to 10",
            "over 10",
        ]
    )
    # One example block -- four grade rows under a single size range,
    # column A merged down them to show the "blank means carried down
    # from above" convention. Deliberately no values in any of them.
    first_data_row = sheet.max_row + 1
    for grade in _GRADE_EXAMPLES:
        sheet.append(["", grade])
    sheet.cell(row=first_data_row, column=1, value="1 to 6")
    sheet.merge_cells(
        start_row=first_data_row, start_column=1, end_row=first_data_row + len(_GRADE_EXAMPLES) - 1, end_column=1
    )

    sheet["A1"].comment = Comment(_INSTRUCTIONS_V2, "StackFit Pro")

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
